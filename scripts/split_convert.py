import sys
import os
import zipfile
import re
import subprocess
import shutil
from openpyxl import load_workbook
from openpyxl.worksheet.properties import WorksheetProperties, PageSetupProperties
from pypdf import PdfWriter

xlsx_src   = sys.argv[1]
output_pdf = sys.argv[2]
lo_path    = sys.argv[3]
tmp_dir    = sys.argv[4]

RANGES = [(1, 59), (60, 100), (101, 141), (142, 189)]


def make_page(src, start_row, end_row, dst):
    wb = load_workbook(src)
    ws = wb.worksheets[0]
    total = ws.max_row
    if end_row < total:
        ws.delete_rows(end_row + 1, total - end_row)
    if start_row > 1:
        ws.delete_rows(1, start_row - 1)
    ws.page_setup.paperSize   = 9
    ws.page_setup.orientation = 'portrait'
    ws.sheet_properties = WorksheetProperties()
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth  = 1
    ws.page_setup.fitToHeight = 1
    wb.save(dst)

    with zipfile.ZipFile(dst, 'r') as z:
        files = {n: z.read(n) for n in z.namelist()}
    xml = files['xl/worksheets/sheet1.xml'].decode('utf-8')
    xml = xml.replace('<pageSetUpPr/>', '<pageSetUpPr fitToPage="1"/>')
    xml = re.sub(
        r'<pageSetup[^/]*/>',
        '<pageSetup paperSize="9" orientation="portrait" fitToWidth="1" fitToHeight="1"/>',
        xml
    )
    files['xl/worksheets/sheet1.xml'] = xml.encode('utf-8')
    with zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)


pdf_files = []
for i, (s, e) in enumerate(RANGES):
    page_xlsx = os.path.join(tmp_dir, f'page_{i}.xlsx')
    make_page(xlsx_src, s, e, page_xlsx)

    profile = os.path.join(tmp_dir, f'lo_profile_{i}')
    profile_url = 'file:///' + profile.replace(os.sep, '/').lstrip('/')
    subprocess.run(
        [lo_path, '--headless', '--norestore',
         f'-env:UserInstallation={profile_url}',
         '--convert-to', 'pdf', page_xlsx, '--outdir', tmp_dir],
        check=True,
        capture_output=True,
        timeout=60
    )

    pdf = page_xlsx.replace('.xlsx', '.pdf')
    if os.path.exists(pdf):
        pdf_files.append(pdf)

    try:
        shutil.rmtree(profile)
    except Exception:
        pass

if not pdf_files:
    print('ERROR: no se generaron PDFs', file=sys.stderr)
    sys.exit(1)

writer = PdfWriter()
for p in pdf_files:
    writer.append(p)
with open(output_pdf, 'wb') as f:
    writer.write(f)

print(f'OK: {len(pdf_files)} paginas -> {output_pdf}')
